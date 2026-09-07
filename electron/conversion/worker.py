import io
import importlib.util
import json
import math
import re
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

from PIL import Image, ImageCms, ImageOps, ImageSequence
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import html
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ConversionError(Exception):
    def __init__(self, message, code="ERR_UTILITY_CONVERSION"):
        super().__init__(message)
        self.code = code


class NeedsOcr(Exception):
    pass


def send(kind, **values):
    print("CREWFLOW_JSON:" + json.dumps({"type": kind, **values}, ensure_ascii=False), flush=True)


def progress(message, percent=55, indeterminate=False):
    send("progress", message=message, percent=percent, indeterminate=indeterminate)


def clean_text(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", "" if value is None else str(value))


def set_font(run, name="Arial", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color and re.fullmatch(r"[0-9A-Fa-f]{6}", color):
        run.font.color.rgb = RGBColor.from_string(color)


def new_document():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.05
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return doc


def set_page(section, width=595.3, height=841.9, margin=36):
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    section.left_margin = section.right_margin = Pt(margin)
    section.top_margin = section.bottom_margin = Pt(margin)
    section.header_distance = section.footer_distance = Pt(12)


def finish_docx(doc, output):
    doc.save(output)
    verified = Document(output)
    if not verified.paragraphs and not verified.tables:
        raise ConversionError("没有生成有效的 Word 内容", "ERR_UTILITY_OUTPUT")


def cell_property(cell, tag, **attrs):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn(f"w:{tag}"))
    if node is None:
        node = OxmlElement(f"w:{tag}")
        props.append(node)
    for key, value in attrs.items():
        node.set(qn(f"w:{key}"), str(value))
    return node


def make_table(doc, data, widths=None, merges=(), header=True):
    if not data or not any(data):
        return None
    rows, columns = len(data), max(map(len, data))
    if rows * columns > 100000:
        raise ConversionError("表格过大，请缩小工作表范围后重试")
    table = doc.add_table(rows=rows, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    available = (doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin) / 12700
    widths = widths or [1] * columns
    scale = available / max(1, sum(widths))
    for index, column in enumerate(table.columns):
        column.width = Pt(max(5, widths[index] * scale))
    for row in table.rows:
        for col, cell in enumerate(row.cells):
            cell.width = Pt(max(5, widths[col] * scale))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_property(cell, "tcMar")
    for r, values in enumerate(data):
        for c, value in enumerate(values):
            run = table.cell(r, c).paragraphs[0].add_run(clean_text(value))
            set_font(run, size=10, bold=header and r == 0)
    for r1, c1, r2, c2 in merges:
        if 0 <= r1 <= r2 < rows and 0 <= c1 <= c2 < columns:
            merged = table.cell(r1, c1).merge(table.cell(r2, c2))
            for paragraph in list(merged.paragraphs)[1:]:
                if not paragraph.text:
                    paragraph._element.getparent().remove(paragraph._element)
    if header:
        row_props = table.rows[0]._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def xlsx_color(value):
    if value and value.type == "rgb" and isinstance(value.rgb, str):
        return value.rgb[-6:]
    return None


def display_excel(value, number_format):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        has_time = isinstance(value, datetime) and any((value.hour, value.minute, value.second))
        return value.strftime("%Y-%m-%d %H:%M:%S" if has_time else "%Y-%m-%d")
    if isinstance(value, (int, float)):
        fmt = number_format.split(";")[0]
        if "%" in fmt:
            digits = len(re.search(r"\.([0#]+)", fmt).group(1)) if re.search(r"\.([0#]+)", fmt) else 0
            return f"{value * 100:.{digits}f}%"
        match = re.search(r"[0#](?:,[0#]+)*\.([0#]+)", fmt)
        if match:
            result = format(value, f",.{len(match.group(1))}f" if "," in fmt else f".{len(match.group(1))}f")
            symbol = next((s for s in ("¥", "￥", "$", "€", "£") if s in fmt), "")
            return symbol + result
        if re.fullmatch(r"0{2,}", fmt):
            return str(int(value)).zfill(len(fmt))
        return str(int(value)) if math.isfinite(value) and int(value) == value else str(value)
    return clean_text(value)


def excel_axis_size(sheet, index, column=False):
    if column:
        dimension = sheet.column_dimensions.get(get_column_letter(index))
        width = dimension.width if dimension and dimension.width else sheet.sheet_format.defaultColWidth or 13
        return width * 5.25 + 3.75
    dimension = sheet.row_dimensions.get(index)
    return (dimension.height if dimension and dimension.height is not None else sheet.sheet_format.defaultRowHeight) or 15


def excel_image_position(sheet, picture):
    anchor = picture.anchor
    marker = getattr(anchor, "_from", None)
    if marker is not None:
        row, col = marker.row + 1, marker.col + 1
        x, y = marker.colOff / 12700, marker.rowOff / 12700
    elif hasattr(anchor, "pos"):
        row = col = 1
        x, y = anchor.pos.x / 12700, anchor.pos.y / 12700
    else:
        raise ConversionError("无法读取表格图片的位置，请重新插入图片后重试")
    width, height = picture.width * 0.75, picture.height * 0.75
    if hasattr(anchor, "ext"):
        width, height = anchor.ext.cx / 12700, anchor.ext.cy / 12700
    elif marker is not None and hasattr(anchor, "to"):
        width = sum(excel_axis_size(sheet, c, True) for c in range(marker.col + 1, anchor.to.col + 1))
        width += (anchor.to.colOff - marker.colOff) / 12700
        height = sum(excel_axis_size(sheet, r) for r in range(marker.row + 1, anchor.to.row + 1))
        height += (anchor.to.rowOff - marker.rowOff) / 12700

    def containing_cell(index, offset, extent, column=False):
        best_index, overlap = index, 0
        limit = 16384 if column else 1048576
        while index <= limit and offset + extent > 0:
            size = excel_axis_size(sheet, index, column)
            covered = max(0, min(size, offset + extent) - max(0, offset))
            if covered > overlap:
                best_index, overlap = index, covered
            offset -= size
            index += 1
        return best_index

    row = containing_cell(row, y, max(1, height))
    col = containing_cell(col, x, max(1, width), True)
    return row, col, max(1, width), y, x


def add_excel_picture(cell, picture, width, max_height):
    data = picture._data()
    crop = getattr(getattr(getattr(picture.anchor, "pic", None), "blipFill", None), "srcRect", None)
    with Image.open(io.BytesIO(data)) as original:
        image = ImageOps.exif_transpose(original)
        if crop is not None:
            left, top, right, bottom = [max(0, min(0.99, (getattr(crop, edge, None) or 0) / 100000)) for edge in ("l", "t", "r", "b")]
            box = (round(image.width * left), round(image.height * top), round(image.width * (1 - right)), round(image.height * (1 - bottom)))
            if box[2] > box[0] and box[3] > box[1]:
                image = image.crop(box)
        if crop is not None or original.format not in ("PNG", "JPEG", "GIF", "BMP", "TIFF") or image.size != original.size:
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            data = buffer.getvalue()
        ratio = image.height / image.width
    width = min(width, max_height / ratio)
    paragraph = cell.add_paragraph() if cell.paragraphs[-1].text or cell.paragraphs[-1]._p.xpath(".//w:drawing") else cell.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(io.BytesIO(data), width=Pt(width), height=Pt(width * ratio))


def excel_to_word(source, output, original_source=None):
    values_book = load_workbook(source, data_only=True)
    styles_book = load_workbook(original_source or source, data_only=False)
    doc = new_document()
    written = 0
    try:
        for sheet in styles_book:
            if sheet.sheet_state != "visible":
                continue
            pictures = [(excel_image_position(sheet, picture), picture) for picture in sheet._images]
            pictures = [(position, picture) for position, picture in pictures if not getattr(sheet.row_dimensions.get(position[0]), "hidden", False) and not getattr(sheet.column_dimensions.get(get_column_letter(position[1])), "hidden", False)]
            populated = [cell for row in sheet for cell in row if cell.value is not None]
            if not populated and not pictures:
                continue
            max_row = max([cell.row for cell in populated] + [position[0] for position, _ in pictures])
            max_col = max([cell.column for cell in populated] + [position[1] for position, _ in pictures])
            for merged in sheet.merged_cells.ranges:
                if merged.min_row <= max_row and merged.min_col <= max_col:
                    max_row = max(max_row, merged.max_row)
                    max_col = max(max_col, merged.max_col)
            row_indexes = [r for r in range(1, max_row + 1) if not sheet.row_dimensions[r].hidden]
            col_indexes = [c for c in range(1, max_col + 1) if not getattr(sheet.column_dimensions.get(get_column_letter(c)), "hidden", False)]
            if not row_indexes or not col_indexes:
                continue
            if len(row_indexes) * len(col_indexes) > 100000:
                raise ConversionError("表格范围过大，请先删除多余的空行或空列")
            section = doc.sections[0] if not written else doc.add_section(WD_SECTION_START.NEW_PAGE)
            widths = [max(21, excel_axis_size(sheet, c, True)) for c in col_indexes]
            landscape = sum(widths) > 523
            set_page(section, 841.9 if landscape else 595.3, 595.3 if landscape else 841.9)
            heading = doc.add_paragraph()
            heading.paragraph_format.space_after = Pt(8)
            set_font(heading.add_run(sheet.title), size=13, bold=True)
            data = []
            for r in row_indexes:
                row_values = []
                for c in col_indexes:
                    cell = sheet.cell(r, c)
                    value = values_book[sheet.title].cell(r, c).value
                    if cell.data_type == "f" and value is None:
                        raise ConversionError("部分公式没有计算结果，请在 Excel 中重新计算并保存后重试")
                    row_values.append(display_excel(value, cell.number_format))
                data.append(row_values)
            merges = []
            for merged in sheet.merged_cells.ranges:
                rs = [i for i, r in enumerate(row_indexes) if merged.min_row <= r <= merged.max_row]
                cs = [i for i, c in enumerate(col_indexes) if merged.min_col <= c <= merged.max_col]
                if rs and cs:
                    merges.append((min(rs), min(cs), max(rs), max(cs)))
            table = make_table(doc, data, widths, merges, header=False)
            for ri, r in enumerate(row_indexes):
                for ci, c in enumerate(col_indexes):
                    original = sheet.cell(r, c)
                    if original.__class__.__name__ == "MergedCell":
                        continue
                    cell = table.cell(ri, ci)
                    fill = xlsx_color(original.fill.fgColor) if original.fill.patternType == "solid" else None
                    if fill:
                        cell_property(cell, "shd", fill=fill)
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(original.alignment.horizontal, WD_ALIGN_PARAGRAPH.LEFT)
                        for run in paragraph.runs:
                            set_font(run, name=original.font.name or "Arial", size=min(18, max(9, original.font.sz or 11)), bold=bool(original.font.b), color=xlsx_color(original.font.color))
            row_map = {r: i for i, r in enumerate(row_indexes)}
            col_map = {c: i for i, c in enumerate(col_indexes)}
            available = (section.page_width - section.left_margin - section.right_margin) / 12700
            max_height = (section.page_height - section.top_margin - section.bottom_margin) / 12700 - 64
            for (r, c, width, _, _), picture in sorted(pictures, key=lambda entry: (entry[0][0], entry[0][1], entry[0][3], entry[0][4])):
                cell = table.cell(row_map[r], col_map[c])
                cell_width = cell.width / 12700
                add_excel_picture(cell, picture, max(1, min(width * available / sum(widths), cell_width - 12)), max_height)
            written += 1
            progress(f"正在转换工作表：{sheet.title}", 50 + written / max(1, len(styles_book.worksheets)) * 35)
        if not written:
            raise ConversionError("工作簿里没有可转换的可见表格")
        finish_docx(doc, output)
    finally:
        values_book.close()
        styles_book.close()
    return "已保留表格和插图，公式保留计算结果；图表和条件格式不包含在内"


def write_excel_tables(tables, output):
    if not tables:
        raise ConversionError("没有识别到可转换的表格", "ERR_UTILITY_NO_TABLE")
    book = Workbook()
    book.remove(book.active)
    border = Border(*([Side(style="thin", color="D9D9D9")] * 4))
    for index, model in enumerate(tables):
        data = model["data"]
        if not data:
            continue
        title = re.sub(r"[\\/*?:\[\]]", "_", model.get("name", f"表格{index + 1}"))[:31]
        sheet = book.create_sheet(title=title)
        for r, values in enumerate(data, 1):
            for c, value in enumerate(values, 1):
                text = clean_text(value)
                cell = sheet.cell(r, c, text)
                if re.fullmatch(r"-?(?:0|[1-9]\d{0,13})(?:\.\d+)?", text):
                    cell.value = float(text) if "." in text else int(text)
                    if "." in text:
                        cell.number_format = "0." + "0" * len(text.split(".")[1])
                else:
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
                cell.font = Font(name="Arial", size=11)
        for r1, c1, r2, c2 in model.get("merges", []):
            sheet.merge_cells(start_row=r1 + 1, start_column=c1 + 1, end_row=r2 + 1, end_column=c2 + 1)
        for c in range(1, sheet.max_column + 1):
            length = max((sum(2 if ord(x) > 255 else 1 for x in str(sheet.cell(r, c).value or "")) for r in range(1, sheet.max_row + 1)), default=8)
            sheet.column_dimensions[get_column_letter(c)].width = min(55, max(10, length + 2))
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape" if sheet.max_column > 5 else "portrait"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    if not book.sheetnames:
        raise ConversionError("没有识别到可转换的表格", "ERR_UTILITY_NO_TABLE")
    book.save(output)
    load_workbook(output, read_only=True).close()


def word_to_excel(source, output):
    from docx.table import Table
    doc = Document(source)
    models = []
    for element in doc.element.body.iter(qn("w:tbl")):
        table = Table(element, doc)
        data, merges, active = [], [], {}
        for ri, row in enumerate(table._tbl.tr_lst):
            before = row.trPr.find(qn("w:gridBefore")) if row.trPr is not None else None
            col = int(before.get(qn("w:val"), "0")) if before is not None else 0
            values = [""] * col
            for tc in row.tc_lst:
                props = tc.tcPr
                span_node = props.find(qn("w:gridSpan")) if props is not None else None
                span = int(span_node.get(qn("w:val"), "1")) if span_node is not None else 1
                paragraphs = tc.findall(qn("w:p"))
                text = "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in paragraphs)
                vertical = props.find(qn("w:vMerge")) if props is not None else None
                continuing = vertical is not None and vertical.get(qn("w:val"), "continue") == "continue"
                values.extend(["" if continuing else text] + [""] * (span - 1))
                if continuing and col in active:
                    merges[active[col]][2] = ri
                elif span > 1 or vertical is not None:
                    merges.append([ri, col, ri, col + span - 1])
                    if vertical is not None:
                        active[col] = len(merges) - 1
                else:
                    active.pop(col, None)
                col += span
            data.append(values)
        models.append({"data": data, "merges": merges})
    write_excel_tables(models, output)
    return "已提取 Word 内的表格，每张表格保存在独立工作表中"


def open_image(source):
    if Path(source).suffix.lower() in (".heic", ".heif"):
        from pillow_heif import register_heif_opener
        register_heif_opener()
    image = Image.open(source)
    return image


def image_frames(source):
    with open_image(source) as image:
        for frame in ImageSequence.Iterator(image):
            yield ImageOps.exif_transpose(frame.copy())


def rgb_image(image):
    if image.mode == "CMYK" and image.info.get("icc_profile"):
        try:
            return ImageCms.profileToProfile(image, ImageCms.ImageCmsProfile(io.BytesIO(image.info["icc_profile"])), ImageCms.createProfile("sRGB"), outputMode="RGB")
        except (ValueError, OSError):
            pass
    if image.mode in ("RGBA", "LA") or "transparency" in image.info:
        rgba = image.convert("RGBA")
        background = Image.new("RGB", image.size, "white")
        background.paste(rgba, mask=rgba.getchannel("A"))
        return background
    return image.convert("RGB")


def save_image(image, destination, target):
    if target == "jpg":
        rgb_image(image).save(destination, "JPEG", quality=95, subsampling=0)
    elif target == "webp":
        image.convert("RGBA" if "A" in image.getbands() else "RGB").save(destination, "WEBP", lossless=True)
    elif target == "bmp":
        rgb_image(image).save(destination, "BMP")
    elif target == "tiff":
        image.save(destination, "TIFF", compression="tiff_deflate")
    else:
        image.convert("RGBA" if "A" in image.getbands() else "RGB").save(destination, "PNG")
    with Image.open(destination) as check:
        check.verify()


_ocr = None
_table_ocr = None
TABLE_MODEL_PATH = None


def get_ocr():
    global _ocr
    if _ocr is None:
        if importlib.util.find_spec("rapidocr") is None:
            raise NeedsOcr()
        from rapidocr import RapidOCR
        progress("正在加载本地文字识别模型", 45, True)
        _ocr = RapidOCR(params={"EngineConfig.onnxruntime.intra_op_num_threads": 4, "EngineConfig.onnxruntime.inter_op_num_threads": 1})
    return _ocr


def ocr_image(image):
    engine = get_ocr()
    import numpy as np
    result = engine(np.array(rgb_image(image))[:, :, ::-1])
    if not result.txts:
        raise ConversionError("没有识别到文字，请使用清晰、方向正确的图片", "ERR_UTILITY_NO_TEXT")
    return result


def table_from_html(value):
    root = html.fromstring(value)
    rows, occupied, merges = [], set(), []
    for r, tr in enumerate(root.xpath(".//tr")):
        row = []
        c = 0
        for td in tr.xpath("./td|./th"):
            while (r, c) in occupied:
                row.append("")
                c += 1
            rowspan = max(1, min(10000, int(td.get("rowspan", 1))))
            colspan = max(1, min(1000, int(td.get("colspan", 1))))
            row.extend([td.text_content().strip()] + [""] * (colspan - 1))
            if rowspan > 1 or colspan > 1:
                merges.append([r, c, r + rowspan - 1, c + colspan - 1])
                occupied.update((rr, cc) for rr in range(r, r + rowspan) for cc in range(c, c + colspan))
            c += colspan
        rows.append(row)
    if len(rows) < 2 or max(map(len, rows), default=0) < 2:
        raise ConversionError("未识别到完整表格，请选择包含清晰行列的表格图片", "ERR_UTILITY_NO_TABLE")
    return {"data": rows, "merges": merges}


def recognize_table(image):
    global _table_ocr
    result = ocr_image(image)
    if _table_ocr is None:
        if importlib.util.find_spec("rapid_table") is None:
            raise NeedsOcr()
        from rapid_table import ModelType, RapidTable, RapidTableInput
        progress("正在加载本地表格识别模型", 55, True)
        if not TABLE_MODEL_PATH or not Path(TABLE_MODEL_PATH).exists():
            raise NeedsOcr()
        _table_ocr = RapidTable(RapidTableInput(model_type=ModelType.SLANETPLUS, model_dir_or_path=TABLE_MODEL_PATH,
                                             engine_cfg={"intra_op_num_threads": 4, "inter_op_num_threads": 1},
                                             ocr_params={"EngineConfig.onnxruntime.intra_op_num_threads": 4}))
    import numpy as np
    progress("正在识别表格行列和合并单元格", 65, True)
    table = _table_ocr(np.array(rgb_image(image))[:, :, ::-1], [(result.boxes, result.txts, result.scores)])
    return table_from_html(table.pred_htmls[0])


def ocr_to_docx(image, doc):
    result = ocr_image(image)
    regions = ruled_table_regions(image)
    blocks = []
    accepted = []
    for box in regions:
        count = sum(inside((float(points[:, 0].min()), float(points[:, 1].min()), float(points[:, 0].max()), float(points[:, 1].max())), box) for points in result.boxes)
        if count < 4:
            continue
        try:
            model = recognize_table(image.crop(box))
        except ConversionError as error:
            if error.code == "ERR_UTILITY_NO_TABLE":
                continue
            raise
        blocks.append((box, "table", model))
        accepted.append(box)
    lines = []
    for points, text in sorted(zip(result.boxes, result.txts), key=lambda item: float(item[0][:, 1].min())):
        box = (float(points[:, 0].min()), float(points[:, 1].min()), float(points[:, 0].max()), float(points[:, 1].max()))
        if any(inside(box, region) for region in accepted):
            continue
        center = (box[1] + box[3]) / 2
        if lines and abs(lines[-1][0] - center) < (box[3] - box[1]) * 0.55:
            lines[-1][1].append((box, text))
        else:
            lines.append((center, [(box, text)]))
    for _, items in lines:
        items.sort(key=lambda item: item[0][0])
        box = (items[0][0][0], min(item[0][1] for item in items), items[-1][0][2], max(item[0][3] for item in items))
        blocks.append((box, "text", "  ".join(item[1] for item in items)))
    for box, kind, content in sorted(blocks, key=lambda block: block[0][1]):
        if kind == "table":
            make_table(doc, content["data"], merges=content["merges"], header=False)
        else:
            paragraph = doc.add_paragraph()
            size = max(9, min(20, (box[3] - box[1]) / max(1, image.width / 523) * 0.75))
            set_font(paragraph.add_run(clean_text(content)), size=size)
            paragraph.paragraph_format.space_after = Pt(3)


def ruled_table_regions(image):
    import cv2
    import numpy as np
    gray = cv2.cvtColor(np.array(rgb_image(image)), cv2.COLOR_RGB2GRAY)
    binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 15, 8)
    horizontal = cv2.morphologyEx(binary, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(30, image.width // 25), 1)))
    vertical = cv2.morphologyEx(binary, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(25, image.height // 30))))
    lines = cv2.dilate(cv2.bitwise_or(horizontal, vertical), np.ones((3, 3), dtype=np.uint8))
    contours, _ = cv2.findContours(lines, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    boxes = []
    for contour in contours:
        x, y, width, height = cv2.boundingRect(contour)
        intersections = cv2.bitwise_and(horizontal[y:y+height, x:x+width], vertical[y:y+height, x:x+width])
        junctions = cv2.connectedComponents(intersections)[0] - 1
        if width > image.width * 0.15 and height > 45 and junctions >= 6:
            boxes.append((max(0, x-5), max(0, y-5), min(image.width, x+width+5), min(image.height, y+height+5)))
    return boxes


def image_convert(sources, target, output):
    total_pixels = 0
    for source in sources:
        with open_image(source) as image:
            total_pixels += image.width * image.height * getattr(image, "n_frames", 1)
    if total_pixels > 120_000_000:
        raise ConversionError("图片总尺寸过大，请减少文件数量或先缩小图片后重试")
    frames = [frame for source in sources for frame in image_frames(source)]
    if not frames:
        raise ConversionError("图片没有可读取的画面")
    try:
        if target == "pdf":
            pages = [rgb_image(frame) for frame in frames]
            pages[0].save(output, "PDF", save_all=True, append_images=pages[1:], resolution=150)
        elif target == "docx":
            doc = new_document()
            for index, frame in enumerate(frames):
                if index:
                    doc.add_page_break()
                set_page(doc.sections[-1])
                progress(f"正在识别第 {index + 1} / {len(frames)} 张图片", 50 + index / len(frames) * 35)
                ocr_to_docx(frame, doc)
            finish_docx(doc, output)
        elif target == "txt":
            text = []
            for index, frame in enumerate(frames):
                progress(f"正在识别第 {index + 1} / {len(frames)} 张图片", 50 + index / len(frames) * 35)
                text.append("\n".join(ocr_image(frame).txts))
            output.write_text("\n\n".join(text) + "\n", encoding="utf-8")
        elif target == "xlsx":
            tables = [recognize_table(frame) for frame in frames]
            write_excel_tables(tables, output)
        elif len(frames) == 1:
            save_image(frames[0], output, target)
        else:
            folder = output.parent / "images"
            folder.mkdir(exist_ok=True)
            for index, frame in enumerate(frames):
                save_image(frame, folder / f"page-{index + 1:04d}.{target}", target)
            return folder, True, "多页图片已逐页转换"
        warning = "识别结果可编辑，请核对文字、数字和表格" if target in ("docx", "txt", "xlsx") else ""
        return output, False, warning
    finally:
        for frame in frames:
            frame.close()


def pdf_table_model(table):
    rows = table.rows
    xs = sorted(set(round(v, 2) for cell in table.cells for v in (cell[0], cell[2])))
    ys = sorted(set(round(v, 2) for cell in table.cells for v in (cell[1], cell[3])))
    data = [[""] * max(1, len(xs) - 1) for _ in range(max(1, len(ys) - 1))]
    merges = []
    extracted = table.extract(x_tolerance=2, y_tolerance=3)
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row.cells):
            if cell is None:
                continue
            c1, c2 = xs.index(round(cell[0], 2)), xs.index(round(cell[2], 2)) - 1
            r1, r2 = ys.index(round(cell[1], 2)), ys.index(round(cell[3], 2)) - 1
            data[r1][c1] = extracted[ri][ci] or ""
            if r2 > r1 or c2 > c1:
                merges.append([r1, c1, r2, c2])
    return {"data": data, "merges": merges, "widths": [b - a for a, b in zip(xs, xs[1:])]}


def inside(box, region):
    x = (box[0] + box[2]) / 2
    y = (box[1] + box[3]) / 2
    return region[0] <= x <= region[2] and region[1] <= y <= region[3]


def unreadable_pdf_text(text):
    return bool(re.search(r"\(cid:\d+\)", text)) or text.count("\ufffd") > max(2, len(text) // 50)


def pdf_to_word_page(page, doc, renderer):
    if unreadable_pdf_text(page.extract_text() or ""):
        bitmap = renderer.render(scale=2)
        try:
            ocr_to_docx(bitmap.to_pil(), doc)
        finally:
            bitmap.close()
        return True
    tables = [t for t in page.find_tables() if len(t.rows) >= 2 and len(t.columns) >= 2]
    blocks = []
    for table in tables:
        blocks.append((table.bbox, "table", pdf_table_model(table)))
    chars = [ch for ch in page.chars if not any(inside((ch["x0"], ch["top"], ch["x1"], ch["bottom"]), t.bbox) for t in tables)]
    if not chars and not tables:
        if not page.images and not page.curves:
            doc.add_paragraph()
            return False
        bitmap = renderer.render(scale=2)
        try:
            ocr_to_docx(bitmap.to_pil(), doc)
        finally:
            bitmap.close()
        return True
    filtered = page.filter(lambda obj: obj.get("object_type") != "char" or obj in chars)
    for line in filtered.extract_text_lines(layout=False, return_chars=True, x_tolerance=2, y_tolerance=3):
        line_chars = line.get("chars", [])
        if not line_chars:
            continue
        fragments = []
        current = []
        for ch in line_chars:
            if current and ch["x0"] - current[-1]["x1"] > max(24, ch.get("size", 11) * 2.5):
                fragments.append(current)
                current = []
            current.append(ch)
        if current:
            fragments.append(current)
        for group in fragments:
            box = (group[0]["x0"], min(ch["top"] for ch in group), group[-1]["x1"], max(ch["bottom"] for ch in group))
            blocks.append((box, "text", group))
    for picture in page.images:
        box = (max(0, picture["x0"]), max(0, picture["top"]), min(page.width, picture["x1"]), min(page.height, picture["bottom"]))
        if box[2] - box[0] < 10 or box[3] - box[1] < 10 or any(inside(box, t.bbox) for t in tables):
            continue
        if (box[2] - box[0]) * (box[3] - box[1]) > page.width * page.height * 0.8:
            continue
        blocks.append((box, "image", None))
    blocks.sort(key=lambda block: (round(block[0][1] / 3), block[0][0]))
    available = (doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin) / 12700
    last_bottom = None
    for box, kind, content in blocks:
        if kind == "table":
            make_table(doc, content["data"], content["widths"], content["merges"], header=False)
        elif kind == "image":
            bitmap = renderer.render(scale=2)
            try:
                crop = bitmap.to_pil().crop(tuple(round(v * 2) for v in box))
                stream = io.BytesIO()
                crop.save(stream, "PNG")
                doc.add_picture(stream, width=Pt(min(available, box[2] - box[0])))
            finally:
                bitmap.close()
        else:
            paragraph = doc.add_paragraph()
            size = max(6, min(48, sum(ch.get("size", 11) for ch in content) / len(content)))
            paragraph.paragraph_format.space_before = Pt(min(12, max(0, box[1] - last_bottom - size * 0.2))) if last_bottom is not None else Pt(0)
            previous = None
            for ch in content:
                text = ch["text"]
                if previous is not None and ch["x0"] - previous["x1"] > size * 0.18 and not text.isspace() and not previous["text"].isspace():
                    paragraph.add_run(" ")
                font = ch.get("fontname", "Arial").split("+")[-1]
                run = paragraph.add_run(clean_text(text))
                set_font(run, name=font, size=max(6, min(48, ch.get("size", size))), bold=bool(re.search("bold|black|heavy", font, re.I)))
                previous = ch
        last_bottom = box[3]
    return False


def pdf_convert(source, target, output):
    import pdfplumber
    import pypdfium2 as pdfium
    tables, texts, used_ocr = [], [], False
    doc = new_document() if target == "docx" else None
    images = output.parent / "images"
    if target in ("png", "jpg"):
        images.mkdir(exist_ok=True)
    try:
        renderer_doc = pdfium.PdfDocument(source)
    except Exception as error:
        raise ConversionError("无法读取 PDF，请确认文件未损坏且没有密码保护") from error
    try:
        with pdfplumber.open(source) as pdf:
            if not pdf.pages:
                raise ConversionError("PDF 没有可转换的页面")
            for index, page in enumerate(pdf.pages):
                progress(f"正在处理第 {index + 1} / {len(pdf.pages)} 页", 50 + index / len(pdf.pages) * 38)
                renderer = renderer_doc[index]
                try:
                    if target in ("png", "jpg"):
                        bitmap = renderer.render(scale=2)
                        try:
                            save_image(bitmap.to_pil(), images / f"page-{index + 1:04d}.{target}", target)
                        finally:
                            bitmap.close()
                    elif target == "txt":
                        text = page.extract_text() or ""
                        if unreadable_pdf_text(text) or (not text.strip() and (page.images or page.curves)):
                            bitmap = renderer.render(scale=2)
                            try:
                                text = "\n".join(ocr_image(bitmap.to_pil()).txts)
                            finally:
                                bitmap.close()
                            used_ocr = True
                        texts.append(text)
                    elif target == "xlsx":
                        found = [t for t in page.find_tables() if len(t.rows) >= 2 and len(t.columns) >= 2]
                        if not found and page.chars:
                            found = [t for t in page.find_tables({"vertical_strategy": "text", "horizontal_strategy": "text", "min_words_vertical": 3}) if len(t.columns) >= 2 and len(t.rows) >= 3]
                        if found:
                            for table in found:
                                tables.append({**pdf_table_model(table), "name": f"第{index + 1}页-表{len(tables) + 1}"})
                        elif not page.chars:
                            bitmap = renderer.render(scale=2)
                            try:
                                model = recognize_table(bitmap.to_pil())
                                tables.append({**model, "name": f"第{index + 1}页"})
                                used_ocr = True
                            finally:
                                bitmap.close()
                    else:
                        section = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
                        set_page(section, page.width, page.height, margin=24)
                        used_ocr = pdf_to_word_page(page, doc, renderer) or used_ocr
                finally:
                    renderer.close()
    finally:
        renderer_doc.close()
    if target == "docx":
        finish_docx(doc, output)
    elif target == "xlsx":
        write_excel_tables(tables, output)
    elif target == "txt":
        if not any(text.strip() for text in texts):
            raise ConversionError("没有提取到文字", "ERR_UTILITY_NO_TEXT")
        output.write_text("\n\n".join(texts) + "\n", encoding="utf-8")
    warning = "识别结果可编辑，请核对文字、数字和表格" if used_ocr else ""
    if target == "docx" and not used_ocr:
        warning = "已转换为可编辑 Word，复杂排版可能需要调整"
    if target == "xlsx" and not used_ocr:
        warning = "已提取可编辑表格，PDF 中的数值不包含原始公式"
    return (images, True, warning) if target in ("png", "jpg") else (output, False, warning)


def main():
    global TABLE_MODEL_PATH
    request = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    TABLE_MODEL_PATH = request.get("tableModelPath")
    source = Path(request["sourcePath"])
    target = request["targetFormat"]
    output = Path(request["outputPath"])
    sources = [Path(value) for value in request.get("sourcePaths", [str(source)])]
    ext = source.suffix.lower().lstrip(".")
    if ext == "xlsx" and target == "docx":
        warning = excel_to_word(source, output, request.get("originalSourcePath"))
        result = output, False, warning
    elif ext == "docx" and target == "xlsx":
        warning = word_to_excel(source, output)
        result = output, False, warning
    elif ext == "pdf" and target in ("docx", "xlsx", "txt", "png", "jpg"):
        result = pdf_convert(source, target, output)
    elif ext in ("png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff", "heic", "heif"):
        result = image_convert(sources, target, output)
    else:
        raise ConversionError("不支持的转换格式", "ERR_UTILITY_FORMAT")
    progress("正在检查转换结果", 95)
    destination, is_directory, warning = result
    if not destination.exists() or (not is_directory and not destination.stat().st_size):
        raise ConversionError("转换结果为空", "ERR_UTILITY_OUTPUT")
    send("result", outputPath=str(destination), isDirectory=is_directory, warning=warning)


if __name__ == "__main__":
    warnings.simplefilter("error", Image.DecompressionBombWarning)
    try:
        main()
    except NeedsOcr:
        send("error", code="ERR_UTILITY_NEEDS_OCR", message="需要准备文字识别组件")
        sys.exit(78)
    except Exception as error:
        send("error", code=getattr(error, "code", "ERR_UTILITY_CONVERSION"), message=str(error))
        sys.exit(1)
